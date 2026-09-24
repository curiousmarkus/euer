import csv
import sqlite3
import unittest
from contextlib import closing
from pathlib import Path

from tests.cli_test_base import BaseCLITestCase

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


class CLIReportsTestCase(BaseCLITestCase):
    def test_summary_rc_does_not_net_off_unpaid_input_vat(self):
        self.write_config('[tax]\nmode = "standard"\n')
        self.add_expense(amount="-100", rc="eu")
        result = self.run_cli(["summary", "--year", "2026"], check=True)
        expenses = result.stdout.split("Umsatzsteuer")[0]
        self.assertRegex(expenses, r"Arbeitsmittel .*?-100\.00 EUR")
        self.assertNotIn("Abziehbare Vorsteuer", expenses)

    def test_summary_entertainment_and_open_review(self):
        self.write_config('[tax]\nmode = "standard"\n')
        self.add_expense(amount="-129", category="Bewirtungsaufwendungen", vat="19")
        result = self.run_cli(["summary", "--year", "2026"], check=True)
        self.assertRegex(result.stdout, r"GESAMT Ausgaben .*?-96\.00 EUR")
        self.assertRegex(result.stdout, r"Abziehbare Vorsteuer \(Zeile 58\).*?-19\.00 EUR")
        self.add_expense(amount="-50", vendor="Other", category="Bewirtungsaufwendungen")
        result = self.run_cli(["summary", "--year", "2026"], check=True)
        self.assertIn("unvollständig", result.stdout)
        self.assertNotIn("GEWINN", result.stdout)

    def test_entertainment_csv_roundtrip(self):
        self.write_config('[tax]\nmode = "standard"\n')
        self.run_cli(
            [
                "add",
                "expense",
                "--date",
                "2026-01-15",
                "--vendor",
                "Restaurant",
                "--category",
                "Bewirtungsaufwendungen",
                "--amount",
                "-129",
                "--vat",
                "19",
                "--tip",
                "10",
            ],
            check=True,
        )
        export_dir = self.root / "exports"

        result = self.run_cli(
            ["export", "--year", "2026", "--format", "csv", "--output", str(export_dir)], check=True
        )
        expense_file = next(
            line.split("Exportiert: ", 1)[1]
            for line in result.stdout.splitlines()
            if line.startswith("Exportiert: ")
        )
        self.db_path = self.root / "roundtrip.db"
        self.run_cli(["init"], check=True)
        self.write_config('[tax]\nmode = "small_business"\n')
        self.run_cli(["import", "--file", expense_file, "--format", "csv"], check=True)
        with sqlite3.connect(self.db_path) as conn:
            row = conn.execute(
                "SELECT amount_eur, vat_input, entertainment_tip_eur, entertainment_vat_status FROM expenses"
            ).fetchone()
        self.assertEqual(row, (-129, 19, 10, "deductible"))

    @unittest.skipUnless(load_workbook, "openpyxl ist ein optionales XLSX-Extra")
    def test_entertainment_xlsx_values_are_numeric(self):
        self.write_config('[tax]\nmode = "standard"\n')
        self.run_cli(
            [
                "add",
                "expense",
                "--date",
                "2026-01-15",
                "--vendor",
                "Restaurant",
                "--category",
                "Bewirtungsaufwendungen",
                "--amount",
                "-129",
                "--vat",
                "19",
                "--tip",
                "10",
            ],
            check=True,
        )
        export_dir = self.root / "exports"
        self.run_cli(
            ["export", "--year", "2026", "--format", "xlsx", "--output", str(export_dir)],
            check=True,
        )
        workbook = load_workbook(next(export_dir.glob("*_Ausgaben.xlsx")))
        try:
            sheet = workbook.worksheets[0]
            headers = {cell.value: cell.column for cell in sheet[1]}
            for name, expected in [
                ("Trinkgeld", 10),
                ("Bewirtung Kostenbasis", 110),
                ("Bewirtung abziehbar", 77),
                ("Bewirtung nicht abziehbar", 33),
            ]:
                cell = sheet.cell(2, headers[name])
                self.assertEqual(cell.data_type, "n")
                self.assertEqual(cell.value, expected)
        finally:
            workbook.close()

    def test_vat_report_requires_year_and_exclusive_period(self):
        missing_year = self.run_cli(["vat-report"])
        self.assertNotEqual(missing_year.returncode, 0)
        conflict = self.run_cli(["vat-report", "--year", "2026", "--quarter", "1", "--month", "1"])
        self.assertNotEqual(conflict.returncode, 0)

    def test_vat_report_table_and_csv_export(self):
        self.run_cli(["setup"], input="\n\n\n\n\nstandard\n", check=True)
        self.add_income(amount="1190.00", vat_rate="19")
        self.add_expense(vendor="EU SaaS", amount="-100.00", rc="eu")
        self.add_expense(vendor="Office", amount="-119.00", vat="20.00")

        table = self.run_cli(["vat-report", "--year", "2026", "--quarter", "1"], check=True)
        self.assertIn("USt-Voranmeldung Q1/2026", table.stdout)
        self.assertIn("KZ 81", table.stdout)
        self.assertIn("1.000 EUR", table.stdout)
        self.assertIn("KZ 83", table.stdout)

        export_dir = self.root / "vat-exports"
        result = self.run_cli(
            [
                "vat-report",
                "--year",
                "2026",
                "--quarter",
                "1",
                "--format",
                "csv",
                "--output",
                str(export_dir),
            ],
            check=True,
        )
        exported = [
            Path(line.split("Exportiert: ", 1)[1])
            for line in result.stdout.splitlines()
            if line.startswith("Exportiert: ")
        ]
        self.assertEqual(len(exported), 2)
        rows = list(csv.reader(exported[0].read_text(encoding="utf-8-sig").splitlines()))
        self.assertIn("kennzahl", rows[0])
        kennzahlen = {row[6] for row in rows[1:]}
        self.assertIn("81", kennzahlen)
        self.assertIn("83", kennzahlen)

    def test_summary_and_rc(self):
        self.add_expense(
            vendor="OpenAI",
            category="Laufende EDV-Kosten",
            amount="-100.00",
            rc=True,
        )
        self.add_income(amount="1000.00")
        result = self.run_cli(["summary", "--year", "2026"], check=True)
        self.assertIn("GESAMT Ausgaben", result.stdout)
        self.assertIn("GESAMT Einnahmen", result.stdout)
        self.assertIn("Umsatzsteuer (Kleinunternehmer)", result.stdout)

    def test_summary_warns_for_legacy_rc_without_jurisdiction(self):
        with closing(sqlite3.connect(self.db_path)) as conn, conn:
            conn.execute(
                """INSERT INTO expenses
                   (uuid, payment_date, vendor, amount_eur, rc_type, vat_input,
                    vat_output, hash)
                   VALUES (?, ?, ?, ?, 'unclassified', 0.0, 19.0, ?)""",
                (
                    "legacy-rc",
                    "2026-01-15",
                    "Legacy SaaS",
                    -100.0,
                    "legacy-rc-hash",
                ),
            )

        result = self.run_cli(["summary", "--year", "2026"], check=True)
        self.assertIn("Reverse-Charge-Buchung(en) ohne EU-/Drittland-Typ", result.stdout)
        self.assertIn("--rc eu|third-country", result.stdout)

    def test_summary_include_private(self):
        self.add_private_deposit(amount="250.00", description="Einlage")
        result = self.run_cli(
            ["summary", "--year", "2026", "--include-private"],
            check=True,
        )
        self.assertIn("Privatvorgänge", result.stdout)
        self.assertIn("Privateinlagen (Zeile 108)", result.stdout)


if __name__ == "__main__":
    unittest.main()
