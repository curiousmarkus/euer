import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None


REPO_ROOT = Path(__file__).resolve().parents[1]
CLI = [sys.executable, "-m", "euercli"]


@unittest.skipUnless(openpyxl, "openpyxl ist nicht installiert")
class XlsxExportTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.db_path = self.root / "test.db"
        self.home = self.root / "home"
        self.home.mkdir()
        self.env = os.environ.copy()
        self.env["HOME"] = str(self.home)
        self.env["USERPROFILE"] = str(self.home)
        self.env["APPDATA"] = str(self.home / "AppData" / "Roaming")
        self.env["PYTHONIOENCODING"] = "utf-8"
        self.run_cli(["init"])

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def run_cli(self, args: list[str]) -> subprocess.CompletedProcess[str]:
        result = subprocess.run(
            CLI + ["--db", str(self.db_path)] + args,
            text=True,
            encoding="utf-8",
            capture_output=True,
            cwd=REPO_ROOT,
            env=self.env,
        )
        self.assertEqual(
            result.returncode,
            0,
            msg=f"Command failed: {args}\nstdout:\n{result.stdout}\nstderr:\n{result.stderr}",
        )
        return result

    def test_standard_xlsx_export_contains_transactions(self) -> None:
        self.run_cli(
            [
                "add",
                "expense",
                "--date",
                "2026-01-15",
                "--vendor",
                "Bürobedarf GmbH",
                "--category",
                "Arbeitsmittel",
                "--amount",
                "-119.00",
            ]
        )
        self.run_cli(
            [
                "add",
                "income",
                "--date",
                "2026-01-20",
                "--source",
                "Kunde GmbH",
                "--category",
                "Umsatzsteuerpflichtige Betriebseinnahmen",
                "--amount",
                "1190.00",
            ]
        )

        output_dir = self.root / "exports"
        result = self.run_cli(
            ["export", "--year", "2026", "--format", "xlsx", "--output", str(output_dir)]
        )

        exported = [
            Path(line.removeprefix("Exportiert: "))
            for line in result.stdout.splitlines()
            if line.startswith("Exportiert: ")
        ]
        self.assertEqual(len(exported), 3)
        self.assertTrue(all(path.exists() for path in exported))

        expense_workbook = openpyxl.load_workbook(output_dir / "EÜR_2026_Ausgaben.xlsx")
        expense_sheet = expense_workbook["Ausgaben"]
        self.assertEqual(expense_sheet.cell(1, 4).value, "Lieferant")
        self.assertEqual(expense_sheet.cell(2, 4).value, "Bürobedarf GmbH")
        self.assertEqual(expense_sheet.cell(2, 6).value, -119)
        expense_workbook.close()

        income_workbook = openpyxl.load_workbook(output_dir / "EÜR_2026_Einnahmen.xlsx")
        income_sheet = income_workbook["Einnahmen"]
        self.assertEqual(income_sheet.cell(1, 4).value, "Quelle")
        self.assertEqual(income_sheet.cell(2, 4).value, "Kunde GmbH")
        self.assertEqual(income_sheet.cell(2, 6).value, 1190)
        income_workbook.close()

        private_workbook = openpyxl.load_workbook(output_dir / "EÜR_2026_Privatvorgaenge.xlsx")
        self.assertEqual(private_workbook.sheetnames, ["PrivateTransfers", "Sacheinlagen"])
        private_workbook.close()

    def test_vat_report_xlsx_contains_report_and_diagnostics(self) -> None:
        self.run_cli(["setup", "--set", "tax.mode", "standard"])
        self.run_cli(
            [
                "add",
                "income",
                "--date",
                "2026-03-10",
                "--source",
                "Kunde GmbH",
                "--category",
                "Umsatzsteuerpflichtige Betriebseinnahmen",
                "--amount",
                "1190.00",
                "--vat-rate",
                "19",
            ]
        )

        output_dir = self.root / "vat-exports"
        result = self.run_cli(
            [
                "vat-report",
                "--year",
                "2026",
                "--quarter",
                "1",
                "--format",
                "xlsx",
                "--output",
                str(output_dir),
            ]
        )

        exported = Path(result.stdout.strip().removeprefix("Exportiert: "))
        self.assertTrue(exported.exists())
        workbook = openpyxl.load_workbook(exported, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Kennzahlen", "Diagnose"])

        report_rows = list(workbook["Kennzahlen"].iter_rows(values_only=True))
        self.assertIn("kennzahl", report_rows[0])
        kennzahl_index = report_rows[0].index("kennzahl")
        self.assertIn("81", {row[kennzahl_index] for row in report_rows[1:]})

        diagnostic_header = next(workbook["Diagnose"].iter_rows(values_only=True))
        self.assertIn("status", diagnostic_header)
        workbook.close()
